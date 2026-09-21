"""
Unit and integration test for ATS backend engine.
"""
import sys
import os
import io

# Add backend directory to sys.path
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from skills_taxonomy import extract_skills_from_text, find_category_for_skill
from parser import (
    evaluate_ats_readability,
    extract_contact_info,
    extract_candidate_name,
    parse_resume_document,
    extract_text_and_layout_from_docx
)
from vector_engine import vector_engine
from ats_scorer import analyze_ats_match
from resume_optimizer import optimize_resume
from exporter import create_docx_from_text
import docx

SAMPLE_RESUME = """
Alex Mercer
alex.mercer@email.com | (555) 234-5678 | github.com/alexmercer | linkedin.com/in/alexmercer | Seattle, WA

PROFESSIONAL SUMMARY
Experienced Full Stack Engineer with 5+ years of experience building modern web applications using React, JavaScript, and Python. Passionate about clean code, performance, and user experience.

TECHNICAL SKILLS
• Programming Languages: JavaScript, Python, TypeScript, HTML5, CSS3
• Frameworks: React, Node.js, Express, Tailwind CSS
• Databases: PostgreSQL, MongoDB
• Tools: Git, GitHub, VS Code

WORK EXPERIENCE
Senior Frontend Developer | TechNova Solutions | 2022 - Present
• Spearheaded frontend modernization using React and TypeScript, boosting render speed by 35%.
• Collaborated with backend engineers to integrate RESTful APIs and optimize state management.
• Mentored 4 junior software developers and conducted weekly code reviews.

Software Engineer | Apex Cloud Systems | 2019 - 2022
• Developed modular UI components and responsive dashboards using React and Node.js.
• Integrated PostgreSQL database schemas and optimized client-side caching.
• Worked in an Agile Scrum environment participating in sprint planning and daily standups.

EDUCATION
Bachelor of Science in Computer Science
University of Washington, 2019
"""

SAMPLE_JD = """
Senior Full Stack Engineer
CloudScale Technologies is looking for a Senior Full Stack Engineer with 4+ years of experience.

Key Responsibilities:
- Design and architect scalable cloud microservices using Python, FastAPI, and Docker.
- Build high-performance user interfaces using React, Next.js, and TypeScript.
- Implement robust CI/CD pipelines using GitHub Actions and deploy on AWS.
- Manage relational and NoSQL databases like PostgreSQL and Redis.
- Drive unit testing with PyTest and Jest to maintain 90%+ code coverage.
- Collaborate across cross-functional teams in an Agile environment.

Requirements:
- 4+ years of software development experience.
- Strong proficiency in Python, FastAPI, React, Next.js, Docker, Kubernetes, AWS, Redis.
- Experience with CI/CD, Git, GitHub Actions, and Jest.
- Bachelor's degree in Computer Science or equivalent.
"""

def test_engine():
    print("1. Testing skills extraction...")
    skills = extract_skills_from_text(SAMPLE_RESUME)
    assert "Programming Languages" in skills, "Failed to find Programming Languages"
    print(f"   Extracted categories: {list(skills.keys())}")
    
    print("\n2. Testing ATS readability evaluation...")
    readability = evaluate_ats_readability(SAMPLE_RESUME)
    assert readability["score"] >= 80, f"Readability score lower than expected: {readability['score']}"
    print(f"   Readability score: {readability['score']}%")

    print("\n3. Testing candidate name and contact detection...")
    contacts = extract_contact_info(SAMPLE_RESUME)
    assert contacts["email"] == "alex.mercer@email.com", f"Email mismatch: {contacts['email']}"
    assert contacts["phone"] is not None, "Phone not found"
    assert contacts["location"] is not None, "Location not found"
    name = extract_candidate_name(SAMPLE_RESUME, contacts["email"])
    assert name == "Alex Mercer", f"Expected Alex Mercer, got: {name}"
    print(f"   Candidate Name: {name}, Location: {contacts['location']}")
    
    print("\n4. Testing parse_resume_document...")
    parsed_doc = parse_resume_document(SAMPLE_RESUME.encode('utf-8'), "alex_mercer_resume.txt")
    assert parsed_doc["candidate_name"] == "Alex Mercer"
    assert parsed_doc["word_count"] > 100
    assert "skills" in parsed_doc["sections_detected"]
    print(f"   Layout detected: {parsed_doc['layout_detected']}")
    print(f"   Sections detected: {parsed_doc['sections_detected']}")

    print("\n5. Testing 128-D Dense Embedding Vector Generation...")
    vec_data = vector_engine.generate_document_vector(parsed_doc["text"], parsed_doc["sections"])
    assert vec_data["dimension"] == 128, f"Dimension expected 128, got {vec_data['dimension']}"
    assert len(vec_data["vector"]) == 128, f"Vector length expected 128, got {len(vec_data['vector'])}"
    assert abs(vec_data["stats"]["norm"] - 1.0) < 0.01, f"Norm expected ~1.0, got {vec_data['stats']['norm']}"
    assert vec_data["stats"]["active_dimensions"] > 50, "Too few active dimensions"
    assert len(vec_data["top_dimensions"]) >= 6, "Top dimensions not generated"
    assert len(vec_data["section_embeddings"]) >= 2, "Section embeddings not generated"
    print(f"   Vector Dimension: {vec_data['dimension']}-D")
    print(f"   L2 Norm: {vec_data['stats']['norm']}")
    print(f"   Active Dimensions: {vec_data['stats']['active_dimensions']}")
    print(f"   Top Dimension 0: {vec_data['top_dimensions'][0]}")
    print(f"   Section embeddings created for: {[s['section'] for s in vec_data['section_embeddings']]}")

    print("\n6. Testing DOCX sequential table parsing...")
    test_doc = docx.Document()
    test_doc.add_paragraph("Jane Doe")
    test_doc.add_paragraph("jane.doe@example.com | San Francisco, CA")
    test_doc.add_heading("Technical Skills Table", level=1)
    table = test_doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Frontend"
    table.rows[0].cells[1].text = "React, TypeScript, CSS"
    table.rows[1].cells[0].text = "Backend"
    table.rows[1].cells[1].text = "Python, FastAPI, Docker"
    test_doc.add_paragraph("Experience as Lead Architect at InnoTech.")
    docx_io = io.BytesIO()
    test_doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    docx_parsed = extract_text_and_layout_from_docx(docx_bytes)
    assert "Frontend | React, TypeScript, CSS" in docx_parsed["text"], "Table content missing in docx parsing"
    assert "Experience as Lead Architect" in docx_parsed["text"], "Post-table content missing"
    print(f"   DOCX parsed successfully with layout: {docx_parsed['layout_detected']}")
    
    print("\n7. Testing ATS Match Analysis...")
    match_result = analyze_ats_match(SAMPLE_RESUME, SAMPLE_JD)
    print(f"   Initial ATS Overall Score: {match_result['overall_score']}% ({match_result['grade']})")
    print(f"   Matched Skills ({match_result['keywords']['total_matched']}): {match_result['keywords']['matched_skills']}")
    print(f"   Missing Skills ({match_result['keywords']['total_missing']}): {match_result['keywords']['missing_skills']}")
    assert match_result["overall_score"] > 0, "ATS score is 0"
    assert "resume_vector" in match_result["vector_analysis"], "Resume vector missing in vector_analysis"
    assert "jd_vector" in match_result["vector_analysis"], "JD vector missing in vector_analysis"
    print(f"   Vector Cosine Sim: {match_result['vector_analysis']['overall_vector_similarity']}%")
    
    print("\n8. Testing Resume Optimization / Tailoring...")
    tailored = optimize_resume(SAMPLE_RESUME, SAMPLE_JD, target_role="Senior Full Stack Engineer")
    print(f"   Projected ATS Score: {tailored['projected_score']}% (Delta: +{tailored['score_delta']}%)")
    print(f"   Generated {len(tailored['tailored_bullets'])} STAR bullets.")
    assert tailored["projected_score"] >= match_result["overall_score"], "Optimized score should be >= original score"
    
    print("\nALL BACKEND ENGINE TESTS PASSED SUCCESSFULLY! [OK]")

if __name__ == "__main__":
    test_engine()
